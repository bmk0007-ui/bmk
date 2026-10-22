"""Helpers for rebuilding image-only slides as editable PowerPoint objects.

All coordinates are in pixels of the 960x720 source slide images
(1 px = 9525 EMU, so the 4:3 slide is exactly 10in x 7.5in).
"""
import io
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

PX = 9525
FONT = "맑은 고딕"
SPC = -40  # default tracking (1/100 pt); the source deck uses tight letter spacing
IMG_DIR = os.environ.get("SLIDE_IMG_DIR", "img")

NAVY = "0B2A5B"
NAVY2 = "1F3B70"
BLUE = "1E6FC0"
SKY = "3C8FD6"
LIGHTBG = "EEEEEE"
GRAY = "6E6E6E"
DGRAY = "404040"
WHITE = "FFFFFF"
BLACK = "000000"
TXT = "1A1A1A"

SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "round": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "hex": MSO_SHAPE.HEXAGON,
    "chevron": MSO_SHAPE.CHEVRON,
    "pentagon": MSO_SHAPE.PENTAGON,
    "rarrow": MSO_SHAPE.RIGHT_ARROW,
    "darrow": MSO_SHAPE.DOWN_ARROW,
    "larrow": MSO_SHAPE.LEFT_ARROW,
    "uarrow": MSO_SHAPE.UP_ARROW,
    "tri": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "dtri": MSO_SHAPE.FLOWCHART_MERGE,
    "para": MSO_SHAPE.PARALLELOGRAM,
    "trap": MSO_SHAPE.TRAPEZOID,
    "callout": MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
    "snip": MSO_SHAPE.SNIP_1_RECTANGLE,
    "round2": MSO_SHAPE.ROUND_2_SAME_RECTANGLE,
    "plaque": MSO_SHAPE.PLAQUE,
    "diamond": MSO_SHAPE.DIAMOND,
    "donut": MSO_SHAPE.DONUT,
    "circarrow": MSO_SHAPE.CIRCULAR_ARROW,
    "star": MSO_SHAPE.STAR_5_POINT,
    "cross": MSO_SHAPE.CROSS,
    "rtri": MSO_SHAPE.RIGHT_TRIANGLE,
    "notch": MSO_SHAPE.NOTCHED_RIGHT_ARROW,
}

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}
VALIGN = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def E(v):
    return Emu(int(round(v * PX)))


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def is_dark(fill):
    if fill is None:
        return False
    h = fill[0] if isinstance(fill, tuple) else fill
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return 0.299 * r + 0.587 * g + 0.114 * b < 150


def new_deck():
    prs = Presentation()
    prs.slide_width = E(960)
    prs.slide_height = E(720)
    return prs


class Slide:
    """Wraps a pptx slide; `src` is the source image number for crops."""

    def __init__(self, prs, num):
        self.prs = prs
        self.num = num
        self.s = prs.slides.add_slide(prs.slide_layouts[6])
        self.src = os.path.join(IMG_DIR, "s%03d.png" % num)
        self._img = None

    # ---------- low level ----------
    def _fill(self, shape, fill):
        if fill is None:
            shape.fill.background()
        elif isinstance(fill, tuple):  # gradient (c1, c2, angle)
            c1, c2 = fill[0], fill[1]
            ang = fill[2] if len(fill) > 2 else 0
            shape.fill.gradient()
            shape.fill.gradient_angle = ang
            st = shape.fill.gradient_stops
            st[0].color.rgb = rgb(c1)
            st[0].position = 0
            st[1].color.rgb = rgb(c2)
            st[1].position = 1.0
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)

    def _line(self, shape, line, lw=1.0, dash=None):
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(lw)
            if dash:
                from pptx.enum.dml import MSO_LINE_DASH_STYLE as D
                shape.line.dash_style = {"dash": D.DASH, "dot": D.ROUND_DOT, "sysdash": D.SQUARE_DOT}[dash]

    def _text(self, tf, text, size=11, color=TXT, bold=False, align="c", valign="m",
              margin=(3, 2, 3, 2), spacing=None, wrap=True, italic=False, font=FONT):
        tf.word_wrap = wrap
        tf.auto_size = None
        l, t, r, b = margin
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = E(l), E(t), E(r), E(b)
        tf.vertical_anchor = VALIGN[valign]
        paras = text if isinstance(text, list) else str(text).split("\n")
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            popts = {}
            runs = p
            if isinstance(p, dict):  # {"runs": [...], "align": .., "size": .., "space": ..}
                popts = p
                runs = p.get("runs", p.get("t", ""))
            if isinstance(runs, str):
                runs = [(runs, {})]
            elif isinstance(runs, tuple):
                runs = [runs]
            para.alignment = ALIGN[popts.get("align", align)]
            sp = popts.get("spacing", spacing)
            if sp:
                para.line_spacing = sp
            if popts.get("before"):
                para.space_before = Pt(popts["before"])
            for rt in runs:
                if isinstance(rt, str):
                    rt = (rt, {})
                txt, o = rt
                run = para.add_run()
                run.text = txt
                f = run.font
                f.name = o.get("font", font)
                f.size = Pt(o.get("s", popts.get("size", size)))
                f.bold = o.get("b", popts.get("bold", bold))
                f.italic = o.get("i", italic)
                if o.get("u"):
                    f.underline = True
                f.color.rgb = rgb(o.get("c", popts.get("color", color)))
                rPr = run._r.get_or_add_rPr()
                rPr.set("spc", str(int(o.get("spc", popts.get("spc", SPC)))))
                ea = rPr.find(qn("a:ea"))
                if ea is None:
                    ea = etree.SubElement(rPr, qn("a:ea"))
                ea.set("typeface", f.name)
        return tf

    # ---------- shapes ----------
    def box(self, x, y, w, h, fill=None, line=None, shape="rect", text=None, lw=1.0,
            radius=None, dash=None, rot=0, shadow=False, adj=None, flipH=False, **tk):
        sp = self.s.shapes.add_shape(SHAPES.get(shape, shape), E(x), E(y), E(w), E(h))
        self._fill(sp, fill)
        self._line(sp, line, lw, dash)
        if radius is not None and shape in ("round", "callout", "round2", "snip"):
            sp.adjustments[0] = radius
        if adj:
            for i, v in enumerate(adj):
                if v is not None:
                    sp.adjustments[i] = v
        if rot:
            sp.rotation = rot
        if flipH:
            sp._element.spPr.find(qn("a:xfrm")).set("flipH", "1")
        if not shadow:
            etree.SubElement(sp._element.spPr, qn("a:effectLst"))
            er = sp._element.find(".//" + qn("a:effectRef"))
            if er is not None:
                er.set("idx", "0")
        if text is not None:
            tk.setdefault("color", WHITE if is_dark(fill) else TXT)
            self._text(sp.text_frame, text, **tk)
        return sp

    def text(self, x, y, w, h, text, size=11, color=TXT, align="l", valign="t", **tk):
        tb = self.s.shapes.add_textbox(E(x), E(y), E(w), E(h))
        self._text(tb.text_frame, text, size=size, color=color, align=align, valign=valign, **tk)
        return tb

    def line(self, x1, y1, x2, y2, color=GRAY, lw=1.0, dash=None, head=None, tail=None):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE as D
            c.line.dash_style = {"dash": D.DASH, "dot": D.ROUND_DOT, "sysdash": D.SQUARE_DOT}[dash]
        ln = c.line._get_or_add_ln()
        if head:
            etree.SubElement(ln, qn("a:headEnd")).set("type", head)
        if tail:
            etree.SubElement(ln, qn("a:tailEnd")).set("type", tail)
        return c

    def arrow(self, x1, y1, x2, y2, color=GRAY, lw=1.25, **k):
        return self.line(x1, y1, x2, y2, color=color, lw=lw, tail="triangle", **k)

    def crop(self, x, y, w, h, dx=None, dy=None, dw=None, dh=None):
        """Insert the source-image region (x,y,w,h) as a picture (photos, logos, complex art)."""
        if self._img is None:
            self._img = Image.open(self.src).convert("RGB")
        im = self._img.crop((x, y, x + w, y + h))
        scale = 3
        im = im.resize((w * scale, h * scale), Image.LANCZOS) if max(w, h) < 400 else im
        bio = io.BytesIO()
        im.save(bio, "PNG")
        bio.seek(0)
        return self.s.shapes.add_picture(bio, E(dx if dx is not None else x), E(dy if dy is not None else y),
                                         E(dw or w), E(dh or h))

    def table(self, x, y, w, h, rows, colw=None, rowh=None, size=9, header_fill=NAVY, header_color=WHITE,
              fill=WHITE, line="A6A6A6", align="c", bold_header=True, fills=None, colors=None,
              bolds=None, merges=None, aligns=None, lw=0.75, header_rows=1, valign="m", sizes=None):
        nr, nc = len(rows), max(len(r) for r in rows)
        gt = self.s.shapes.add_table(nr, nc, E(x), E(y), E(w), E(h))
        tbl = gt.table
        # strip default style banding
        tblPr = tbl._tbl.tblPr
        tblPr.set("firstRow", "0")
        tblPr.set("bandRow", "0")
        if colw:
            tot = sum(colw)
            for i, cw in enumerate(colw):
                tbl.columns[i].width = E(w * cw / tot)
        if rowh:
            tot = sum(rowh)
            for i, rh in enumerate(rowh):
                tbl.rows[i].height = E(h * rh / tot)
        else:
            for i in range(nr):
                tbl.rows[i].height = E(h / nr)
        for m in merges or []:
            r1, c1, r2, c2 = m
            tbl.cell(r1, c1).merge(tbl.cell(r2, c2))
        for r in range(nr):
            for c in range(nc):
                cell = tbl.cell(r, c)
                val = rows[r][c] if c < len(rows[r]) else ""
                if val is None:
                    continue
                hdr = r < header_rows
                f = (fills or {}).get((r, c), (fills or {}).get(r, header_fill if hdr else fill))
                if f is None:
                    cell.fill.background()
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(f)
                col = (colors or {}).get((r, c), (colors or {}).get(r, header_color if hdr else TXT))
                b = (bolds or {}).get((r, c), (bolds or {}).get(r, bold_header if hdr else False))
                a = (aligns or {}).get((r, c), (aligns or {}).get(c, align)) if not hdr else "c"
                sz = (sizes or {}).get((r, c), (sizes or {}).get(r, size))
                cell.text_frame.paragraphs[0].text = ""
                self._text(cell.text_frame, val, size=sz, color=col, bold=b, align=a, valign=valign,
                           margin=(3, 1, 3, 1))
                cell.vertical_anchor = VALIGN[valign]
                tcPr = cell._tc.get_or_add_tcPr()
                for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                    ln = etree.SubElement(tcPr, qn(tag), w=str(int(lw * 12700)))
                    if line is None:
                        etree.SubElement(ln, qn("a:noFill"))
                    else:
                        sf = etree.SubElement(ln, qn("a:solidFill"))
                        etree.SubElement(sf, qn("a:srgbClr"), val=line)
                # lnX must precede fill in schema: move fill elements to the end
                for tag in ("a:solidFill", "a:noFill"):
                    for el in tcPr.findall(qn(tag)):
                        tcPr.remove(el)
                        tcPr.append(el)
        return gt

    # ---------- recurring template pieces ----------
    def band(self, x, y, w, h, cut=10, fill=NAVY):
        """Navy band with the bottom-right corner cut diagonally."""
        sp = self.box(x, y, w, h, fill=fill, shape="snip", adj=[cut / h])
        sp._element.spPr.find(qn("a:xfrm")).set("flipV", "1")
        return sp

    def header(self, num, title, sub=None, sub_bold=None, tag=None, sub_w=510, title_w=None, sub_size=14):
        """Section header used on content slides: shield logo, navy title band with
        number box, right 'BU VISION 2030', optional sub-title tab and right tag."""
        tw = title_w or (78 + 27 * len(title))
        self.box(95, 67, 830, 3, fill=NAVY)
        self.band(95, 22, tw, 48, cut=11)
        self.crop(28, 12, 70, 76)  # shield logo
        self.box(111, 31, 30, 30, fill=NAVY2, line=WHITE, lw=1.25, text=str(num), size=17, bold=True,
                 color=WHITE, margin=(0, 0, 0, 0))
        self.text(148, 22, tw - 55, 46, title, size=22, bold=True, color=WHITE, valign="m")
        self.text(770, 50, 155, 20, [[("BU VISION ", {"c": NAVY}), ("2030", {"c": SKY})]],
                  size=10.5, bold=True, align="r", valign="m", margin=(0, 0, 0, 0))
        if sub:
            self.band(32, 97, sub_w, 37, cut=10)
            runs = [(sub_bold, {"b": True}), (sub, {"b": False})] if sub_bold else [(sub, {"b": True})]
            self.text(40, 97, sub_w - 10, 37, [runs], size=sub_size, color=WHITE, valign="m")
        if tag:
            self.box(805, 108, 120, 20, fill="F3F6FA", line=None)
            self.text(805, 108, 118, 20, tag, size=9, color=SKY, align="r", valign="m", margin=(0, 0, 2, 0))

    def panel(self, x=35, y=134, w=890, h=566, fill="EFEFEF", line="9DB3D0"):
        return self.box(x, y, w, h, fill=fill, line=line, lw=1.0)

    def notes(self, txt):
        self.s.notes_slide.notes_text_frame.text = txt


# ---------- reusable compound parts ----------
def num_callout(S, x, y, w, h, num, text, size=11.5, color=NAVY, line=NAVY2):
    """White rounded callout (tail to the right) with a numbered circle on top."""
    S.box(x, y, w, h, fill=WHITE, line=line, lw=1.5, shape="callout", adj=[0.56, -0.12, 0.16667],
          text=text, size=size, bold=True, color=color, spacing=1.1)
    cx = x + w / 2
    S.box(cx - 15, y - 12, 30, 24, fill=WHITE, line=line, lw=1.5, shape="round", radius=0.3,
          text=str(num), size=12, bold=True, color=color, margin=(0, 0, 0, 0))


def detail_box(S, x, y, w, h, head, items, size=11, head_size=None, line="8EA3BF", tx=None):
    """Rounded white box with a '▶ head' line and '→ item' lines."""
    S.box(x, y, w, h, fill=WHITE, line=line, lw=1.0, shape="round", radius=0.08)
    paras = [[("▶ " + head, {"b": True, "c": NAVY, "s": head_size or size})]]
    paras += [[("→ " + it, {"c": DGRAY})] for it in items]
    S.text(tx if tx is not None else x + 85, y, w - (tx - x if tx else 85) - 5, h, paras, size=size,
           valign="m", spacing=1.15)


def arrow_paras(body, mark="→ ", indent="    "):
    """body: str or list of str; each str becomes '→ first line' plus indented continuation lines."""
    out = []
    for b in ([body] if isinstance(body, str) else body):
        lines = b.split("\n")
        out.append([(mark + lines[0], {})])
        out += [[(indent + ln, {})] for ln in lines[1:]]
    return out


def card_items(S, x, y, w, h, items, size=11.5, gap=12, valign="t", spacing=1.12):
    paras = []
    for i, (head, body) in enumerate(items):
        paras.append({"runs": [(head, {"b": True, "c": NAVY})], "before": 0 if i == 0 else gap})
        paras += arrow_paras(body)
    return S.text(x, y, w, h, paras, size=size, color=TXT, valign=valign, spacing=spacing)


def env_card(S, x, w, label, num, items, concl, y=160, h=455, cy=628, ch=57, size=11.5, label_w=137,
             concl_align="c", gap=12, concl_size=11, valign="t"):
    """Analysis card: gray label tab, numbered badge, bold-heading + '→' items, navy conclusion box."""
    S.box(x, y, w, h, fill=WHITE, line="D0D7E2", shape="round", radius=0.04)
    S.box(x + 3, y - 10, label_w, 30, fill=("6B6B6B", "A0A0A0", 0), shape="snip", adj=[0.5],
          text=label, size=13, bold=True, color=WHITE, align="l", margin=(14, 0, 0, 0))
    S.box(x + w - 30, y - 3, 27, 26, fill=WHITE, line=NAVY2, lw=1.5, shape="round", radius=0.2,
          text=str(num), size=13, bold=True, color=NAVY2, margin=(0, 0, 0, 0))
    paras = []
    for i, (head, body) in enumerate(items):
        paras.append({"runs": [(head, {"b": True, "c": NAVY})], "before": 0 if i == 0 else gap})
        paras += arrow_paras(body)
    S.text(x + 22, y + 22, w - 30, h - 30, paras, size=size, color=TXT, valign=valign, spacing=1.12)
    S.box(x + w / 2 - 20, cy - 12, 40, 12, fill=NAVY, shape="dtri")
    S.box(x, cy, w, ch, fill=("1F4E8C", "0B2A5B", 90), text=concl, size=concl_size, bold=False, color=WHITE,
          spacing=1.05, align=concl_align, margin=(14 if concl_align == "c" else 30, 0, 8, 0))


def source(S, txt, y=702, size=6.5):
    S.text(30, y, 900, 16, txt, size=size, color="7F7F7F", valign="m", margin=(2, 0, 0, 0))


def bar_chart(S, x, y, w, h, cats, vals, vmin=None, vmax=None, major=None, numfmt="General", color="3C9BE0",
              label_size=8, axis_size=8, gridlines=True, gap=150, labels=True, series="값", line_chart=False,
              horizontal=False, label_color="404040", axis_fmt=None):
    """Native, editable column chart (or line/bar)."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(series, vals)
    kind = XL_CHART_TYPE.LINE_MARKERS if line_chart else (
        XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED)
    gf = S.s.shapes.add_chart(kind, E(x), E(y), E(w), E(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(axis_size)
    ch.font.name = FONT
    va = ch.value_axis
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    if major:
        va.major_unit = major
    va.has_major_gridlines = gridlines
    if gridlines:
        va.major_gridlines.format.line.color.rgb = rgb("D9D9D9")
    va.tick_labels.number_format = axis_fmt or numfmt
    va.tick_labels.number_format_is_linked = False
    va.format.line.fill.background()
    ca = ch.category_axis
    ca.format.line.color.rgb = rgb("BFBFBF")
    ca.tick_labels.font.size = Pt(axis_size)
    va.tick_labels.font.size = Pt(axis_size)
    plot = ch.plots[0]
    if not line_chart:
        plot.gap_width = gap
    ser = plot.series[0]
    if line_chart:
        ser.format.line.color.rgb = rgb(color)
    else:
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(color)
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(label_size)
        dl.font.color.rgb = rgb(label_color)
        dl.number_format = numfmt
        dl.number_format_is_linked = False
        if not line_chart:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
    return ch


def fill_plot(ch, x=0.02, y=0.02, w=0.96, h=0.96):
    """Give the chart's plot area a manual layout so it fills the graphic frame."""
    pa = ch._chartSpace.find(".//" + qn("c:plotArea"))
    lay = pa.find(qn("c:layout"))
    if lay is None:
        lay = etree.Element(qn("c:layout"))
        pa.insert(0, lay)
    for c in list(lay):
        lay.remove(c)
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    for tag, val in (("c:layoutTarget", "inner"), ("c:xMode", "edge"), ("c:yMode", "edge"),
                     ("c:x", x), ("c:y", y), ("c:w", w), ("c:h", h)):
        etree.SubElement(ml, qn(tag)).set("val", str(val))
