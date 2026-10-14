"""
교사(校舍)시설 확보율 예측 시뮬레이션 생성 스크립트 (v4 기준)
실행: pip install openpyxl  →  python build_simulation.py
출력: 교사시설확보율_예측시뮬레이션.xlsx (엑셀에서 열면 수식 자동 계산)

[구조]
- YEARS / YC          : 연도 목록과 입력 시트 열(2024·2025 실적, 2026~ 예측)
- build_base(...)     : 연도별 기초자료 시트 생성 (집계표 수치는 bld24/bld25, ent, cur, yr 리스트에 입력)
- 간편계산기          : 재학생 수만 넣으면 예상 재학생기준 확보율이 바로 나오는 첫 화면 (건물 면적은 2025 실적 고정)
- 입력 / 결과요약 / 민감도 / 사용안내 : 각 '# ====' 구간
[새 연도 실적 추가 방법]
1) base2026 시트 생성 + build_base(base2026, 2026, ...) 호출 (집계표 1-3, 2-2, 2-3, 공시값)
2) YEARS 에 실적 연도 추가, 입력 시트의 실적 열 연결과 결과요약의 인정면적 분기(y==2024) 로직 확장
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

FN="맑은 고딕"
def F(bold=False,color="000000",size=10,italic=False): return Font(name=FN,bold=bold,color=color,size=size,italic=italic)
def fill(c): return PatternFill("solid",fgColor=c)
NAVY="1F3864"; BLUE="0000FF"; GREEN="008000"; GRAY="595959"
YEL=fill("FFF2CC"); HDR=fill("1F3864"); SUB=fill("D9E1F2"); TOT=fill("EDEDED"); WHITE=fill("FFFFFF")
GOOD=fill("C6EFCE"); WARN=fill("FFEB9C"); BAD=fill("FFC7CE"); KPI=fill("F2F7FC")
thin=Side(style="thin",color="BFBFBF"); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
C=Alignment(horizontal="center",vertical="center",wrap_text=True)
L=Alignment(horizontal="left",vertical="center",wrap_text=True)
AREA='#,##0;(#,##0);"-"'; AREA2='#,##0.00;(#,##0.00);"-"'; CNT='#,##0;(#,##0);"-"'; PCT='0.0"%"'; PP='+0.0"%p";-0.0"%p";"-"'
YEARS=[2024,2025,2026,2027,2028,2029,2030]; YC="FGHIJKL"; NY=len(YEARS)

def put(ws,ref,v=None,font=None,fl=None,fmt=None,al=None,border=True):
    c=ws[ref]
    if v is not None: c.value=v
    c.font=font or F()
    if fl: c.fill=fl
    if fmt: c.number_format=fmt
    if al: c.alignment=al
    if border: c.border=BD
    return c
def header(ws,row,cols,vals,height=30):
    for col,v in zip(cols,vals): put(ws,f"{col}{row}",v,F(True,"FFFFFF"),HDR,al=C)
    ws.row_dimensions[row].height=height
def banner(ws,ref,text,span):
    put(ws,ref,text,F(True,"FFFFFF",14),fill(NAVY),al=L,border=False); ws.merge_cells(f"{ref}:{span}")
    ws.row_dimensions[int(''.join(ch for ch in ref if ch.isdigit()))].height=30
def note(ws,ref,text,span=None,color=GRAY,bold=False,size=10):
    put(ws,ref,text,F(bold,color,size),al=L,border=False)
    if span: ws.merge_cells(f"{ref}:{span}")
def step(ws,ref,text,span):
    put(ws,ref,text,F(True,NAVY,11),SUB,al=L); ws.merge_cells(f"{ref}:{span}")
    ws.row_dimensions[int(''.join(ch for ch in ref if ch.isdigit()))].height=22

wb=Workbook()
guide=wb.active; guide.title="사용안내"
quick=wb.create_sheet("간편계산기",0)
inp=wb.create_sheet("입력")
res=wb.create_sheet("결과요약")
sen=wb.create_sheet("민감도")
base25=wb.create_sheet("기초자료_2025"); base24=wb.create_sheet("기초자료_2024")
for ws in wb: ws.sheet_view.showGridLines=False

# ============ 기초자료 (연도별) ============
def build_base(b,year,bld,ent,cur,yr,pub,dig,cnote):
    banner(b,"A1",f"기초자료 · {year} 시설현황 집계표 / {year} 정보공시(14-라)","I1")
    note(b,"A2","이 시트는 원자료 보관용입니다. 수정할 필요가 없습니다. (파란 글씨 = 집계표에서 옮겨 적은 값)","I2")
    step(b,"A4","1. 건물 보유면적 (집계표 1-3, ㎡)","I4")
    header(b,5,"ABCDEFGHI",["캠퍼스","위치","교육기본시설","지원시설","연구시설","부속시설","기타시설","임대시설","교사보유면적\n(집계표)"])
    for i,r in enumerate(bld):
        row=6+i; put(b,f"A{row}",r[0],al=C); put(b,f"B{row}",r[1],al=C)
        for j,col in enumerate("CDEFGHI"): put(b,f"{col}{row}",r[2+j],F(color=BLUE),fmt=AREA2)
    step(b,"A12","2. 확보율 계산에 들어가는 면적 (C 기본 · D 지원 · E 연구)","I12")
    header(b,13,"ABCDEFG",["캠퍼스","기본시설(C)","지원시설(D)","연구시설(E)","계","집계표 교사보유면적","차이"])
    put(b,"H13","어떻게 계산했나",F(True,"FFFFFF"),HDR,al=C); b.merge_cells("H13:I13")
    rows=[("본교(천안)","=C6+H6","=D6","=E6","=I6+I7+I8",cnote),
          ("서울캠퍼스","=C9","=D9+D10","=E9","=I9+I10","교지경계선 내 기본+지원 + 20km 이내 학생기숙사(조이하우스) 4,794.83㎡")]
    for i,r in enumerate(rows):
        row=14+i; put(b,f"A{row}",r[0],al=C)
        for col,v in zip("BCD",r[1:4]): put(b,f"{col}{row}",v,fmt=AREA2)
        put(b,f"E{row}",f"=SUM(B{row}:D{row})",fmt=AREA2); put(b,f"F{row}",r[4],fmt=AREA2); put(b,f"G{row}",f"=E{row}-F{row}",fmt=AREA2)
        put(b,f"H{row}",r[5],F(size=9),al=L); b.merge_cells(f"H{row}:I{row}"); b.row_dimensions[row].height=32
    put(b,"A16","합계",F(True),TOT,al=C)
    for col in "BCDEFG": put(b,f"{col}16",f"=SUM({col}14:{col}15)",F(True),TOT,fmt=AREA2)
    put(b,"H16","",fl=TOT); b.merge_cells("H16:I16")
    step(b,"A18","3. 시설정원 (집계표 2-2, 명)","I18")
    header(b,19,"ABCDEFG",["캠퍼스","과정","인문사회","자연과학","공학","예체능","계"])
    put(b,"A20","입학정원 기준 (수업연한 곱한 값)",F(True),TOT,al=L); b.merge_cells("A20:G20")
    def blk(start,data):
        for i,r in enumerate(data):
            row=start+i; put(b,f"A{row}",r[0],al=C); put(b,f"B{row}",r[1],al=C)
            for j,col in enumerate("CDEF"): put(b,f"{col}{row}",r[2+j],F(color=BLUE),fmt=CNT)
            put(b,f"G{row}",f"=SUM(C{row}:F{row})",fmt=CNT)
    blk(21,ent)
    put(b,"A24",f"재학생 기준 ({year}.4.1.)",F(True),TOT,al=L); b.merge_cells("A24:G24")
    blk(25,cur)
    step(b,"A29","4. 본교(천안) 학부 학년별 재학생 (집계표 2-3, 주간 · 야간 0명)","I29")
    header(b,30,"ABCDEFG",["계열","1학년","2학년","3학년","4학년","계","3번 표와 차이"])
    for i,(nm,*v) in enumerate(yr):
        row=31+i; put(b,f"A{row}",nm,al=C)
        for j,col in enumerate("BCDE"): put(b,f"{col}{row}",v[j],F(color=BLUE),fmt=CNT)
        put(b,f"F{row}",f"=SUM(B{row}:E{row})",fmt=CNT); put(b,f"G{row}",f"=F{row}-{'CDEF'[i]}25",fmt=CNT)
    step(b,"A36",f"5. {year} 정보공시(14-라) 값과 이 시트로 계산한 값 대조","I36")
    header(b,37,"ABCDE",["항목","공시값","계산값(원값)","계산값(공시 표기)","차이"])
    items=[("입학정원 기준면적(A)","=SUM(C21:C23)*12+SUM(D21:F23)*14",0),
           ("재학생 기준면적(B)","=SUM(C25:C27)*12+SUM(D25:F27)*14",0),
           ("기본시설(C)","=B16",0),("지원시설(D)","=C16",0),("연구시설(E)","=D16",0),
           ("부속시설(교지경계선 내)","=F6+F9",0),("기타시설(교지경계선 내)","=G6+G9",0),
           ("입학정원기준 확보율(%)","=E16/C38*100",dig),("재학생기준 확보율(%)","=E16/C39*100",dig)]
    rfmt="0" if dig==0 else "0."+"0"*dig
    for i,((lab,fm,d),pv) in enumerate(zip(items,pub)):
        row=38+i; put(b,f"A{row}",lab,al=L)
        put(b,f"B{row}",pv,F(color=BLUE),fmt=CNT if i<7 else rfmt,al=C)
        put(b,f"C{row}",fm,fmt=AREA2 if i<7 else "0.000",al=C)
        put(b,f"D{row}",f"=ROUND(C{row},{d})",fmt=CNT if i<7 else rfmt,al=C)
        put(b,f"E{row}",f"=D{row}-B{row}",F(True),fmt='0.0;-0.0;"일치"',al=C)
    note(b,"A47",f"※ {year} 공시의 확보율은 "+("소수 첫째자리" if dig==1 else "정수")+"로 표기되어 같은 자릿수로 반올림해 비교했습니다. 차이 열이 모두 '일치'면 원자료 입력이 정확한 것입니다.","I47",size=9)
    for col,w in zip("ABCDEFGHI",[22,14,14,15,13,13,15,22,40]): b.column_dimensions[col].width=w

bld25=[("본교(천안)","교지경계선 내",146250.18,49778.99,0,4252.54,497,5178.83,201208),
       ("본교(천안)","20km 이내",6569.73,0,0,2374.25,7778.51,0,0),
       ("본교(천안)","20km 외",0,0,2942.06,0,0,0,0),
       ("서울캠퍼스","교지경계선 내",18570.5,945.96,0,0,0,0,19516.46),
       ("서울캠퍼스","20km 이내",0,4794.83,0,0,0,0,4794.83)]
bld24=[("본교(천안)","교지경계선 내",146250.18,49778.99,0,4252.54,497,5178.83,201208),
       ("본교(천안)","20km 이내",6569.73,0,0,3509.25,16254.57,5904.37,0),
       ("본교(천안)","20km 외",0,0,2942.06,0,0,0,0),
       ("서울캠퍼스","교지경계선 내",18570.5,945.96,0,0,0,0,19516.46),
       ("서울캠퍼스","20km 이내",0,4794.83,0,0,0,0,4794.83)]
build_base(base25,2025,bld25,
    [("본교(천안)","학부",4032,4488,0,2640),("본교(천안)","대학원",55,25,6,0),("서울캠퍼스","대학원",886,0,0,0)],
    [("본교(천안)","학부",4792,4045,0,2737),("본교(천안)","대학원",228,30,229,0),("서울캠퍼스","대학원",1001,0,0,0)],
    [("인문사회",1154,1068,1228,1342),("자연과학",1304,901,972,868),("공학",0,0,0,0),("예체능",694,615,684,744)],
    [159902,170826,170000,55520,0,4253,497,141,132],0,
    "교지경계선 내 기본+지원+연구 + 임대시설(학생후생복지) 5,178.83㎡. 20km 이내 기본 6,569.73㎡ · 20km 외 연구 2,942.06㎡은 미인정")
build_base(base24,2024,bld24,
    [("본교(천안)","학부",4648,3852,0,2660),("본교(천안)","대학원",0,25,20,0),("서울캠퍼스","대학원",900,0,0,0)],
    [("본교(천안)","학부",5106,3707,0,2773),("본교(천안)","대학원",125,20,126,0),("서울캠퍼스","대학원",995,0,0,0)],
    [("인문사회",1274,1188,1221,1423),("자연과학",1091,910,808,898),("공학",0,0,0,0),("예체능",698,612,722,741)],
    [158374,167476,170000,55520,0,4253,497,142.4,134.7],1,
    "교지경계선 내 기본+지원+연구 + 임대시설(학생후생복지) 5,178.83㎡. 20km 이내 기본 6,569.73㎡ · 일반임대 5,904.37㎡ · 20km 외 연구 2,942.06㎡은 미인정")

# ============ 입력 ============
s=inp
banner(s,"A1","입력 · 노란 칸만 고치면 결과요약이 자동으로 바뀝니다","L1")
note(s,"A2","2026년 이후 칸에는 전년도 값이 미리 들어 있습니다. 바뀔 것으로 예상되는 칸만 새 숫자로 덮어쓰세요.  예) 2027년 천안 대학원 공학 재학생 229 → 300명 예상이면 그 칸에 300 입력. 2024·2025는 실적(회색)이라 고치지 않습니다.","L2",color="C00000")
header(s,4,"ABCDE",["캠퍼스","과정","계열","1인당\n기준면적(㎡)",""])
s.merge_cells("D4:E4")
for col,y in zip(YC,YEARS): put(s,f"{col}4",f"{y}\n{'실적' if y<=2025 else '예측'}",F(True,"FFFFFF"),HDR,al=C)
s.row_dimensions[4].height=32
groups=[("본교(천안)","학부",0),("본교(천안)","대학원",1),("서울캠퍼스","대학원",2)]
flds=[("인문사회","C",12),("자연과학","D",14),("공학","E",14),("예체능","F",14)]
def section(start,title,brow0,help_):
    step(s,f"A{start}",title,f"L{start}")
    note(s,f"A{start+1}",help_,f"L{start+1}")
    row=start+2; ranges=[]; subs=[]
    for camp,proc,gi in groups:
        g0=row
        for fld,bcol,unit in flds:
            put(s,f"A{row}",camp if row==g0 else None,al=C); put(s,f"B{row}",proc if row==g0 else None,al=C)
            put(s,f"C{row}",fld,al=C); put(s,f"D{row}",unit,F(color=GRAY),al=C); s.merge_cells(f"D{row}:E{row}")
            put(s,f"F{row}",f"=기초자료_2024!{bcol}{brow0+gi}",F(color=GREEN),fl=fill("F2F2F2"),fmt=CNT)
            put(s,f"G{row}",f"=기초자료_2025!{bcol}{brow0+gi}",F(color=GREEN),fl=fill("F2F2F2"),fmt=CNT)
            for i,col in enumerate(YC[2:]): put(s,f"{col}{row}",f"={YC[i+1]}{row}",fl=YEL,fmt=CNT)
            row+=1
        s.merge_cells(f"A{g0}:A{row-1}"); s.merge_cells(f"B{g0}:B{row-1}")
        ranges.append((g0,row-1))
        put(s,f"A{row}",f"{camp} {proc} 소계",F(True),TOT,al=C); s.merge_cells(f"A{row}:E{row}")
        for col in YC: put(s,f"{col}{row}",f"=SUM({col}{g0}:{col}{row-1})",F(True),TOT,fmt=CNT)
        subs.append(row); row+=1
    put(s,f"A{row}","합계",F(True,"FFFFFF"),fill("44546A"),al=C); s.merge_cells(f"A{row}:E{row}")
    for col in YC: put(s,f"{col}{row}","="+"+".join(f"{col}{r}" for r in subs),F(True,"FFFFFF"),fill("44546A"),fmt=CNT)
    return ranges,row
CUR,CUR_T=section(6,"STEP 1.  재학생 수 (명) — 매년 4월 1일 기준 예상 인원  ★ 가장 중요",25,
    "정원내+정원외 포함 / 계약학과, 특성화고 재직자 특별전형, 학사학위취득유예생은 빼고 입력. 대학원은 주간 재학생.")
ENT,ENT_T=section(CUR_T+2,"STEP 2.  입학정원 시설정원 (명) — 정원 조정이 있을 때만 수정",21,
    "학부: 계열별 입학정원 × 수업연한(4년) / 대학원: 주간 입학정원 × 수업연한. 모집유보 인원, 첨단학과 결손인원 활용분은 제외.")
r=ENT_T+2
step(s,f"A{r}","STEP 3.  교사 면적 증감 (㎡) — 신축·철거·용도변경이 있을 때만 입력 (늘면 +, 줄면 −)",f"L{r}")
note(s,f"A{r+1}","그 해 4월 1일 이전에 준공·멸실된 면적을 해당 연도 칸에 입력. 한 번 넣으면 이후 연도에도 계속 반영됩니다. (2024·2025 실적 면적은 기초자료 시트 값 사용)",f"L{r+1}")
INC0=r+2; incs=[("본교(천안)","기본시설(강의실·실험실·도서관 등)"),("본교(천안)","지원시설(기숙사·강당 등)"),("본교(천안)","연구시설"),
               ("서울캠퍼스","기본시설(강의실·실험실·도서관 등)"),("서울캠퍼스","지원시설(기숙사·강당 등)"),("서울캠퍼스","연구시설")]
for i,(camp,k) in enumerate(incs):
    row=INC0+i
    put(s,f"A{row}",camp if i%3==0 else None,al=C); put(s,f"B{row}",k,al=L); s.merge_cells(f"B{row}:E{row}")
    put(s,f"F{row}",0,F(color=GRAY),fill("F2F2F2"),fmt=AREA); put(s,f"G{row}",0,F(color=GRAY),fill("F2F2F2"),fmt=AREA)
    for col in YC[2:]: put(s,f"{col}{row}",0,F(color=BLUE),YEL,fmt=AREA)
s.merge_cells(f"A{INC0}:A{INC0+2}"); s.merge_cells(f"A{INC0+3}:A{INC0+5}")
INC1=INC0+5
for col,w in zip("ABCDEFGHIJKL",[14,9,11,6,6,11,11,11,11,11,11,11]): s.column_dimensions[col].width=w
s.freeze_panes="F5"

def sp(col,rngs):   # area formula across groups
    return "+".join(f"SUMPRODUCT(입력!$D${a}:$D${z},입력!{col}{a}:{col}{z})" for a,z in rngs)

# ============ 결과요약 ============
R=res
banner(R,"A1","결과요약 · 교사(校舍)시설 확보율 예측 (백석대학교 본교 = 천안 + 서울, 정보공시 기준)","L1")
# KPI area
put(R,"A3","조회 연도 ▶",F(True,NAVY,12),al=C,border=False)
put(R,"B3",2027,F(True,"C00000",14),YEL,fmt="0",al=C)
dv=DataValidation(type="list",formula1='"2024,2025,2026,2027,2028,2029,2030"',allow_blank=False); R.add_data_validation(dv); dv.add("B3")
note(R,"C3","← 칸을 눌러 연도를 고르세요","E3")
# yearly table first (rows 16-21) so KPI can reference
T0=17
step(R,f"A{T0-2}","연도별 한눈에 보기","L"+str(T0-2))
header(R,T0-1,"ABCDEFGHIJKL",["연도","재학생 수\n(명)","재학생\n기준면적(B)","입학정원\n기준면적(A)","교사 인정면적\n(C+D+E)","재학생기준\n확보율","입학정원기준\n확보율","전년 대비\n(재학생기준)","판정","여유 면적\n(인정면적−B)","추가로 받을 수 있는\n재학생(인문사회)","추가로 받을 수 있는\n재학생(자연·공·예체능)"],height=42)
for i,(y,col) in enumerate(zip(YEARS,YC)):
    row=T0+i
    put(R,f"A{row}",y,F(True),fmt="0",al=C)
    put(R,f"B{row}",f"=입력!{col}{CUR_T}",fmt=CNT)
    put(R,f"C{row}","="+sp(col,CUR),fmt=AREA)
    put(R,f"D{row}","="+sp(col,ENT),fmt=AREA)
    put(R,f"E{row}","=기초자료_2024!$E$16" if y==2024 else f"=기초자료_2025!$E$16+SUM(입력!$G${INC0}:{col}{INC1})",fmt=AREA)
    put(R,f"F{row}",f"=IF(C{row}=0,0,E{row}/C{row}*100)",F(True),fmt=PCT,al=C)
    put(R,f"G{row}",f"=IF(D{row}=0,0,E{row}/D{row}*100)",fmt=PCT,al=C)
    put(R,f"H{row}","-" if i==0 else f"=F{row}-F{row-1}",fmt=PP,al=C)
    put(R,f"I{row}",f'=IF(MIN(F{row},G{row})<100,"미달",IF(MIN(F{row},G{row})<사용안내!$C$27,"주의","적정"))',F(True),al=C)
    put(R,f"J{row}",f"=E{row}-C{row}",fmt=AREA)
    put(R,f"K{row}",f"=IF(J{row}<0,0,ROUNDDOWN(J{row}/12,0))",fmt=CNT,al=C)
    put(R,f"L{row}",f"=IF(J{row}<0,0,ROUNDDOWN(J{row}/14,0))",fmt=CNT,al=C)
TL=T0+NY-1
rng=f"I{T0}:I{TL}"
R.conditional_formatting.add(rng,FormulaRule(formula=[f'I{T0}="적정"'],fill=GOOD,font=Font(name=FN,bold=True,color="006100")))
R.conditional_formatting.add(rng,FormulaRule(formula=[f'I{T0}="주의"'],fill=WARN,font=Font(name=FN,bold=True,color="9C5700")))
R.conditional_formatting.add(rng,FormulaRule(formula=[f'I{T0}="미달"'],fill=BAD,font=Font(name=FN,bold=True,color="9C0006")))
R.conditional_formatting.add(f"A{T0}:L{TL}",FormulaRule(formula=[f"$A{T0}=$B$3"],border=Border(top=Side(style="medium",color="C00000"),bottom=Side(style="medium",color="C00000"))))
note(R,f"A{TL+1}","※ 2024·2025는 실제 공시값과 같습니다(2024 재학생기준 134.7%·입학정원기준 142.4% / 2025 132%·141%).  판정: 100% 미만 = 미달, 100% 이상 ~ '주의 기준' 미만 = 주의(주의 기준은 사용안내 시트에서 변경)","L"+str(TL+1),size=9)
R.row_dimensions[TL+1].height=28

idx=f"MATCH($B$3,$A${T0}:$A${TL},0)"
def pick(col): return f"INDEX(${col}${T0}:${col}${TL},{idx})"
cards=[("A5","재학생기준 확보율",pick("F"),PCT),("D5","입학정원기준 확보율",pick("G"),PCT),
       ("G5","여유 면적(㎡)",pick("J"),AREA),("J5","판정",pick("I"),None)]
for ref,lab,fm,fmt in cards:
    col=ref[0]; nc=chr(ord(col)+2)
    put(R,f"{col}5",lab,F(True,"FFFFFF",10),fill("2F5597"),al=C); R.merge_cells(f"{col}5:{nc}5")
    put(R,f"{col}6","="+fm,F(True,NAVY,22),KPI,fmt=fmt,al=C); R.merge_cells(f"{col}6:{nc}7")
R.row_dimensions[6].height=26; R.row_dimensions[7].height=20
R.conditional_formatting.add("J6",FormulaRule(formula=['$J$6="적정"'],fill=GOOD,font=Font(name=FN,bold=True,size=22,color="006100")))
R.conditional_formatting.add("J6",FormulaRule(formula=['$J$6="주의"'],fill=WARN,font=Font(name=FN,bold=True,size=22,color="9C5700")))
R.conditional_formatting.add("J6",FormulaRule(formula=['$J$6="미달"'],fill=BAD,font=Font(name=FN,bold=True,size=22,color="9C0006")))
# plain sentences
sent1=(f'=$B$3&"년 재학생은 "&TEXT({pick("B")},"#,##0")&"명으로, 재학생기준 교사확보율은 "&TEXT({pick("F")},"0.0")&"%입니다. "'
       f'&IF({pick("J")}>=0,"법정 기준(100%)보다 "&TEXT({pick("J")},"#,##0")&"㎡ 여유가 있습니다.","법정 기준(100%)에 "&TEXT(-{pick("J")},"#,##0")&"㎡ 부족합니다.")')
sent2=(f'="→ 이 상태에서 재학생이 100명 늘면 확보율이 약 "&TEXT({pick("F")}-{pick("E")}/({pick("C")}+1200)*100,"0.0")&"%p(인문사회) ~ "'
       f'&TEXT({pick("F")}-{pick("E")}/({pick("C")}+1400)*100,"0.0")&"%p(자연·공학·예체능) 내려갑니다."')
sent3=(f'="→ 확보율 100%를 지키면서 더 받을 수 있는 재학생은 인문사회 기준 약 "&TEXT({pick("K")},"#,##0")&"명, 자연·공학·예체능 기준 약 "&TEXT({pick("L")},"#,##0")&"명입니다."')
for i,fm in enumerate([sent1,sent2,sent3]):
    put(R,f"A{9+i}",fm,F(i==0,NAVY if i==0 else "404040",11 if i==0 else 10),fill("FFFFFF"),al=L,border=False); R.merge_cells(f"A{9+i}:L{9+i}")
    R.row_dimensions[9+i].height=20
# 공시양식 미리보기
P0=TL+4
step(R,f"A{P0-1}","정보공시 14-라 양식 미리보기 (단위: ㎡, %)","L"+str(P0-1))
header(R,P0,"ABCDEFGHIJ",["기준연도","입학정원 기준\n(A)","재학생 기준\n(B)","기본시설\n(C)","지원시설\n(D)","연구시설\n(E)","부속시설","기타시설","입학정원기준\n(C+D+E)/A×100","재학생기준\n(C+D+E)/B×100"],height=42)
for i,(y,col) in enumerate(zip(YEARS,YC)):
    row=P0+1+i; t=T0+i
    put(R,f"A{row}",y,F(True),fmt="0",al=C)
    put(R,f"B{row}",f"=ROUND(D{t},0)",fmt=AREA); put(R,f"C{row}",f"=ROUND(C{t},0)",fmt=AREA)
    if y==2024:
        bs="기초자료_2024"
        put(R,f"D{row}",f"=ROUND({bs}!$B$16,0)",fmt=AREA); put(R,f"E{row}",f"=ROUND({bs}!$C$16,0)",fmt=AREA); put(R,f"F{row}",f"=ROUND({bs}!$D$16,0)",fmt=AREA)
    else:
        bs="기초자료_2025"
        put(R,f"D{row}",f"=ROUND({bs}!$B$16+SUM(입력!$G${INC0}:{col}{INC0})+SUM(입력!$G${INC0+3}:{col}{INC0+3}),0)",fmt=AREA)
        put(R,f"E{row}",f"=ROUND({bs}!$C$16+SUM(입력!$G${INC0+1}:{col}{INC0+1})+SUM(입력!$G${INC0+4}:{col}{INC0+4}),0)",fmt=AREA)
        put(R,f"F{row}",f"=ROUND({bs}!$D$16+SUM(입력!$G${INC0+2}:{col}{INC0+2})+SUM(입력!$G${INC0+5}:{col}{INC0+5}),0)",fmt=AREA)
    put(R,f"G{row}",f"=ROUND({bs}!$F$6+{bs}!$F$9,0)",fmt=AREA)
    put(R,f"H{row}",f"=ROUND({bs}!$G$6+{bs}!$G$9,0)",fmt=AREA)
    put(R,f"I{row}",f"=ROUND(G{t},1)",F(True),fmt="0.0",al=C); put(R,f"J{row}",f"=ROUND(F{t},1)",F(True),fmt="0.0",al=C)
PL=P0+NY
note(R,f"A{PL+1}","※ 확보율은 소수 첫째자리로 표시(2024 공시는 소수 첫째자리, 2025 공시는 정수 표기). 부속·기타시설은 확보율과 무관해 2025 이후는 2025 값을 표시. 2024·2025행은 실제 공시값과 일치합니다.","L"+str(PL+1),size=9)
# campus
K0=PL+4
step(R,f"A{K0-1}","캠퍼스별 확보율 (집계표 1-1과 같은 방식)","L"+str(K0-1))
header(R,K0,"ABCDEFG",["연도","천안\n재학생기준","천안\n입학정원기준","서울\n재학생기준","서울\n입학정원기준","천안\n인정면적","서울\n인정면적"],height=36)
for i,(y,col) in enumerate(zip(YEARS,YC)):
    row=K0+1+i
    put(R,f"A{row}",y,F(True),fmt="0",al=C)
    if y==2024:
        put(R,f"F{row}","=기초자료_2024!$E$14",fmt=AREA); put(R,f"G{row}","=기초자료_2024!$E$15",fmt=AREA)
    else:
        put(R,f"F{row}",f"=기초자료_2025!$E$14+SUM(입력!$G${INC0}:{col}{INC0+2})",fmt=AREA)
        put(R,f"G{row}",f"=기초자료_2025!$E$15+SUM(입력!$G${INC0+3}:{col}{INC1})",fmt=AREA)
    put(R,f"B{row}",f"=F{row}/({sp(col,CUR[:2])})*100",fmt=PCT,al=C)
    put(R,f"C{row}",f"=F{row}/({sp(col,ENT[:2])})*100",fmt=PCT,al=C)
    put(R,f"D{row}",f"=G{row}/({sp(col,CUR[2:])})*100",fmt=PCT,al=C)
    put(R,f"E{row}",f"=G{row}/({sp(col,ENT[2:])})*100",fmt=PCT,al=C)
KL=K0+NY
# ---- 2024→2025 실적 변화 요인
V0=KL+3
step(R,f"A{V0}","2024 → 2025 실적 변화 요인 — 확보율이 왜 바뀌었나?","L"+str(V0))
header(R,V0+1,"ABCDEFG",["계열","1인당\n기준면적","2024\n재학생","2025\n재학생","재학생\n증감","기준면적\n변화(㎡)","확보율\n영향(%p)"],height=36)
fl4=[("인문사회",12,0),("자연과학",14,1),("공학",14,2),("예체능",14,3)]
for i,(nm,u,k) in enumerate(fl4):
    r=V0+2+i
    put(R,f"A{r}",nm,al=C); put(R,f"B{r}",u,fmt='0"㎡"',al=C)
    put(R,f"C{r}","="+"+".join(f"입력!F{a+k}" for a,z in CUR),fmt=CNT)
    put(R,f"D{r}","="+"+".join(f"입력!G{a+k}" for a,z in CUR),fmt=CNT)
    put(R,f"E{r}",f"=D{r}-C{r}",fmt='+#,##0;-#,##0;"-"',al=C)
    put(R,f"F{r}",f"=E{r}*B{r}",fmt='+#,##0;-#,##0;"-"')
    put(R,f"G{r}",f"=IF(F{r}=0,0,$E${T0+1}/($C${T0}+F{r})*100-$F${T0})",fmt=PP,al=C)
rt=V0+6
put(R,f"A{rt}","합계",F(True),TOT,al=C); put(R,f"B{rt}","",fl=TOT)
for col in "CDEF": put(R,f"{col}{rt}",f"=SUM({col}{V0+2}:{col}{V0+5})",F(True),TOT,fmt=CNT if col in "CD" else '+#,##0;-#,##0;"-"',al=C if col=="E" else None)
put(R,f"G{rt}",f"=$F${T0+1}-$F${T0}",F(True),TOT,fmt=PP,al=C)
put(R,f"A{rt+1}","교사 인정면적 변화(㎡)",F(True),SUB,al=C); R.merge_cells(f"A{rt+1}:D{rt+1}")
put(R,f"E{rt+1}",f"=E{T0+1}-E{T0}",F(True),fmt='+#,##0.00;-#,##0.00;"변동 없음"',al=C); R.merge_cells(f"E{rt+1}:G{rt+1}")
put(R,f"A{rt+2}","입학정원 기준면적 변화(㎡)",F(True),SUB,al=C); R.merge_cells(f"A{rt+2}:D{rt+2}")
put(R,f"E{rt+2}",f"=D{T0+1}-D{T0}",F(True),fmt='+#,##0;-#,##0;"변동 없음"',al=C); R.merge_cells(f"E{rt+2}:G{rt+2}")
sent=(f'="▶ 2024 → 2025: 재학생 "&TEXT(E{rt},"+#,##0;-#,##0;0")&"명, 재학생 기준면적 "&TEXT(F{rt},"+#,##0;-#,##0;0")&"㎡, 교사 인정면적 "'
      f'&IF(E{rt+1}=0,"변동 없음",TEXT(E{rt+1},"+#,##0;-#,##0")&"㎡")&" → 재학생기준 확보율 "&TEXT(F{T0},"0.0")&"% → "&TEXT(F{T0+1},"0.0")&"% ("&TEXT(G{rt},"+0.0;-0.0;0.0")&"%p). '
      f'입학정원기준은 "&TEXT(G{T0},"0.0")&"% → "&TEXT(G{T0+1},"0.0")&"%."')
put(R,f"I{V0+1}",sent,F(True,NAVY,10),fill("F2F7FC"),al=Alignment(horizontal="left",vertical="top",wrap_text=True)); R.merge_cells(f"I{V0+1}:L{V0+5}")
put(R,f"I{V0+6}","※ 계열별 '확보율 영향'은 2024 기준면적에 그 계열 변화만 더했을 때의 변화로, 계열별 합이 전체 변화와 정확히 같지는 않습니다(비율 계산 특성).",F(size=8,color=GRAY),al=Alignment(horizontal="left",vertical="top",wrap_text=True),border=False)
R.merge_cells(f"I{V0+6}:L{V0+8}")
KL=rt+3
# chart helper: 100% line
put(R,f"N{T0-1}","기준선",F(color=GRAY),border=False)
for i in range(NY): put(R,f"N{T0+i}",100,F(color="BFBFBF"),border=False)
ch=LineChart(); ch.title="연도별 교사확보율 추이 (%)"; ch.height=8; ch.width=17
ch.y_axis.title="확보율(%)"; ch.y_axis.majorGridlines=None
data=Reference(R,min_col=6,max_col=7,min_row=T0-1,max_row=TL); ch.add_data(data,titles_from_data=True)
ch.add_data(Reference(R,min_col=14,min_row=T0-1,max_row=TL),titles_from_data=True)
ch.set_categories(Reference(R,min_col=1,min_row=T0,max_row=TL))
ch.series[0].graphicalProperties.line.solidFill="2F5597"; ch.series[0].graphicalProperties.line.width=28000
ch.series[1].graphicalProperties.line.solidFill="A5A5A5"; ch.series[1].graphicalProperties.line.width=22000
ch.series[2].graphicalProperties.line.solidFill="C00000"; ch.series[2].graphicalProperties.line.dashStyle="dash"
ch.series[0].dLbls=DataLabelList(); ch.series[0].dLbls.showVal=True; ch.series[0].dLbls.numFmt='0.0'
for k in ("showSerName","showCatName","showLegendKey","showPercent"): setattr(ch.series[0].dLbls,k,False)
ch.series[0].dLbls.position="t"
for sr,nm in zip(ch.series,["재학생기준 확보율","입학정원기준 확보율","법정 기준 100%"]): sr.tx=SeriesLabel(v=nm)
ch.y_axis.scaling.min=80; ch.x_axis.delete=False; ch.y_axis.delete=False
ch.legend.position="b"
ch.visible_cells_only=False
R.add_chart(ch,f"A{KL+3}")
bc=BarChart(); bc.type="col"; bc.title="재학생 기준면적 vs 교사 인정면적 (㎡)"; bc.height=8; bc.width=17
bc.add_data(Reference(R,min_col=3,min_row=T0-1,max_row=TL),titles_from_data=True)
bc.add_data(Reference(R,min_col=5,min_row=T0-1,max_row=TL),titles_from_data=True)
bc.set_categories(Reference(R,min_col=1,min_row=T0,max_row=TL))
bc.series[0].graphicalProperties.solidFill="F4B183"; bc.series[1].graphicalProperties.solidFill="2F5597"
bc.legend.position="b"; bc.y_axis.majorGridlines=None; bc.x_axis.delete=False; bc.y_axis.delete=False
for sr,nm in zip(bc.series,["재학생 기준면적(B)","교사 인정면적(C+D+E)"]): sr.tx=SeriesLabel(v=nm)
R.add_chart(bc,f"A{KL+20}")
for col,w in zip("ABCDEFGHIJKLN",[11,12,13,13,14,12,12,11,9,13,15,15,6]): R.column_dimensions[col].width=w
R.column_dimensions["I"].width=13
R.column_dimensions["N"].hidden=True

# ============ 민감도 ============
m=sen
banner(m,"A1","민감도 분석 · 재학생이 늘거나 줄면 교사확보율은 어떻게 될까?","K1")
note(m,"A2","이 시트로 알 수 있는 것  ①재학생 증감에 따른 확보율 변화  ②확보율 목표를 지키려면 재학생을 최대 몇 명까지 늘릴 수 있는지  ③계열을 섞어 늘릴 때의 결과.  노란 칸만 바꾸면 됩니다.","K2",color="C00000")
m.row_dimensions[2].height=32
# --- 0. 기준
step(m,"A4","0단계 · 기준 정하기 — 어느 해 상태에서 출발할지 고릅니다","K4")
put(m,"A5","기준 연도",F(True),SUB,al=C); put(m,"B5",2025,F(True,"C00000",12),YEL,fmt="0",al=C)
dv2=DataValidation(type="list",formula1='"2024,2025,2026,2027,2028,2029,2030"'); m.add_data_validation(dv2); dv2.add("B5")
put(m,"A6","표 증감 간격(명)",F(True),SUB,al=C); put(m,"B6",250,F(True,BLUE),YEL,fmt=CNT,al=C)
note(m,"A7","↑ 연도는 결과요약 시트의 예측값을 그대로 가져옵니다.","C7",size=9)
iy=f"MATCH($B$5,결과요약!$A${T0}:$A${TL},0)"
base_items=[("D5","기준 재학생 수(명)",f"=INDEX(결과요약!$B${T0}:$B${TL},{iy})",CNT),
            ("D6","재학생 기준면적 B(㎡)",f"=INDEX(결과요약!$C${T0}:$C${TL},{iy})",AREA),
            ("D7","교사 인정면적(㎡)",f"=INDEX(결과요약!$E${T0}:$E${TL},{iy})",AREA),
            ("H5","현재 재학생기준 확보율",f"=IF(E6=0,0,E7/E6*100)",PCT),
            ("H6","여유 면적(㎡)","=E7-E6",AREA),
            ("H7","입학정원 시설정원(명)",f"=INDEX(입력!$F${ENT_T}:$K${ENT_T},{iy})",CNT)]
for ref,lab,fm,fmt in base_items:
    col=ref[0]; row=ref[1:]; vc=chr(ord(col)+1) if col=="D" else chr(ord(col)+2)
    put(m,ref,lab,F(True),SUB,al=C)
    if col=="D": m.merge_cells(f"D{row}:D{row}")
    else: m.merge_cells(f"H{row}:I{row}")
    put(m,f"{vc}{row}",fm,F(True,NAVY,11),KPI,fmt=fmt,al=C)
    if col=="H": m.merge_cells(f"J{row}:K{row}")
# --- 1. 핵심 문장
step(m,"A9","1단계 · 한눈에 보는 핵심 (숫자가 바뀌면 문장도 자동으로 바뀝니다)","K9")
TH="사용안내!$C$27"
sents=[
 '="● 현재("&$B$5&"년) 재학생은 "&TEXT(E5,"#,##0")&"명, 재학생기준 확보율은 "&TEXT(J5,"0.0")&"%이며, 법정 기준(100%)보다 "&TEXT(J6,"#,##0")&"㎡ 여유가 있습니다."',
 '="● 재학생 1명이 늘면 기준면적이 인문사회는 12㎡, 자연과학·공학·예체능은 14㎡ 늘어납니다. 100명 늘면 확보율이 약 "&TEXT(J5-E7/(E6+1200)*100,"0.0")&"%p(인문사회) ~ "&TEXT(J5-E7/(E6+1400)*100,"0.0")&"%p(자연·공학·예체능) 내려갑니다."',
 '=IF(J6<0,"● 이미 법정 기준(100%)에 미달합니다.","● 확보율이 100% 아래로 떨어지려면 재학생이 인문사회만 늘 경우 "&TEXT(ROUNDDOWN(J6/12,0)+1,"#,##0")&"명 이상, 자연·공학·예체능만 늘 경우 "&TEXT(ROUNDDOWN(J6/14,0)+1,"#,##0")&"명 이상 늘어야 합니다.")',
 f'=IF(J5<{TH},"● 이미 주의 기준("&{TH}&"%)보다 낮습니다.","● 주의 기준("&{TH}&"%)까지는 인문사회 "&TEXT(ROUNDDOWN((E7/{TH}*100-E6)/12,0),"#,##0")&"명, 자연·공학·예체능 "&TEXT(ROUNDDOWN((E7/{TH}*100-E6)/14,0),"#,##0")&"명까지 늘어도 됩니다.")',
 '=IF(E5>J7,"● 재학생이 "&TEXT(E5-J7,"#,##0")&"명 넘게 줄면 재학생 수가 입학정원 시설정원("&TEXT(J7,"#,##0")&"명)보다 적어집니다. 이때부터는 재학생 기준면적 이상만 보유하면 법정 기준을 충족합니다(지침 ※ 조항).","● 재학생 수가 이미 입학정원 시설정원보다 적어, 재학생 기준면적이 법정 의무 기준입니다.")']
for i,fm in enumerate(sents):
    r=10+i; put(m,f"A{r}",fm,F(i==0,NAVY if i==0 else "262626",10),al=L,border=False); m.merge_cells(f"A{r}:K{r}"); m.row_dimensions[r].height=30
# --- 2. 원리
step(m,"A16","2단계 · 계산 원리 — 예시 숫자로 따라가 보기","K16")
note(m,"A17","공식:  재학생기준 확보율 = 교사 인정면적 ÷ ( 현재 재학생 기준면적 + 늘어난 재학생 × 1인당 기준면적 ) × 100","K17",color=NAVY,bold=True,size=11)
note(m,"A18","1인당 기준면적은 인문사회 12㎡, 자연과학·공학·예체능 14㎡ (운영 중인 대학, 학부·대학원 구분 없음). 줄어드는 경우는 음수(−)로 넣으면 됩니다.","K18",size=9)
put(m,"A19","예시 증감 인원(명)",F(True),SUB,al=C); m.merge_cells("A19:B19"); put(m,"C19",500,F(True,BLUE,12),YEL,fmt='+#,##0;-#,##0;0',al=C)
put(m,"D19","계열 선택",F(True),SUB,al=C); put(m,"E19","인문사회",F(True,BLUE,11),YEL,al=C); m.merge_cells("E19:F19")
dv3=DataValidation(type="list",formula1='"인문사회,자연·공학·예체능"'); m.add_data_validation(dv3); dv3.add("E19")
put(m,"G19","1인당 기준면적",F(True),SUB,al=C); put(m,"H19",'=IF(E19="인문사회",12,14)',F(True),fmt='0"㎡"',al=C)
ex=[("①","늘어나는 기준면적",'=TEXT(C19,"#,##0")&"명 × "&H19&"㎡"',"=C19*H19",AREA),
    ("②","새 재학생 기준면적",'=TEXT(E6,"#,##0")&" + "&TEXT(C19*H19,"#,##0")',"=E6+C19*H19",AREA),
    ("③","새 확보율",'=TEXT(E7,"#,##0")&" ÷ "&TEXT(E6+C19*H19,"#,##0")&" × 100"',"=IF(E6+C19*H19<=0,0,E7/(E6+C19*H19)*100)",PCT),
    ("④","현재 대비 변화",'=TEXT(J5,"0.0")&"% → "&TEXT(J23,"0.0")&"%"',"=J23-J5",PP)]
header(m,20,"ABCDEFGHIJK",["순서","내용","","계산식","","","","","","결과",""],height=22)
m.merge_cells("B20:C20"); m.merge_cells("D20:I20"); m.merge_cells("J20:K20")
for i,(n,lab,txt,val,fmt) in enumerate(ex):
    r=21+i
    put(m,f"A{r}",n,F(True,NAVY,12),al=C); put(m,f"B{r}",lab,F(True),al=L); m.merge_cells(f"B{r}:C{r}")
    put(m,f"D{r}",txt,al=L); m.merge_cells(f"D{r}:I{r}")
    put(m,f"J{r}",val,F(True,NAVY,11),KPI,fmt=fmt,al=C); m.merge_cells(f"J{r}:K{r}")
# --- 3. 표
S0=27
step(m,f"A{S0}","3단계 · 재학생 증감별 확보율 표 — 한 계열만 늘거나 줄었다고 가정","K"+str(S0))
note(m,f"A{S0+1}","읽는 법: 왼쪽 '재학생 증감'만큼 해당 계열 재학생이 바뀌면 오른쪽 확보율이 됩니다. 파란 줄이 현재. 빨강=100% 미만(미달), 노랑=주의 기준 미만, 초록=적정.","K"+str(S0+1),size=9)
put(m,f"A{S0+2}","",fl=HDR); put(m,f"B{S0+2}","",fl=HDR)
put(m,f"C{S0+2}","인문사회 계열이 바뀔 때 (1인당 12㎡)",F(True,"FFFFFF"),fill("2F5597"),al=C); m.merge_cells(f"C{S0+2}:F{S0+2}")
put(m,f"G{S0+2}","자연과학·공학·예체능이 바뀔 때 (1인당 14㎡)",F(True,"FFFFFF"),fill("548235"),al=C); m.merge_cells(f"G{S0+2}:J{S0+2}")
header(m,S0+3,"ABCDEFGHIJ",["재학생 증감\n(명)","전체 재학생\n(명)","재학생\n기준면적(㎡)","확보율","현재 대비","판정","재학생\n기준면적(㎡)","확보율","현재 대비","판정"],height=34)
D0=S0+4
for i in range(13):
    r=D0+i; k=i-6
    put(m,f"A{r}",f"=$B$6*{k}",F(True),fmt='+#,##0;-#,##0;"0 (현재)"',al=C)
    put(m,f"B{r}",f"=$E$5+A{r}",fmt=CNT,al=C)
    for cA,cR,cD,cJ,u in (("C","D","E","F",12),("G","H","I","J",14)):
        put(m,f"{cA}{r}",f"=$E$6+A{r}*{u}",fmt=AREA)
        put(m,f"{cR}{r}",f"=IF({cA}{r}<=0,0,$E$7/{cA}{r}*100)",F(True),fmt=PCT,al=C)
        put(m,f"{cD}{r}",f"={cR}{r}-$J$5",fmt=PP,al=C)
        put(m,f"{cJ}{r}",f'=IF({cR}{r}<100,"미달",IF({cR}{r}<{TH},"주의","적정"))',F(True),al=C)
DL=D0+12
m.conditional_formatting.add(f"A{D0+6}:J{D0+6}",FormulaRule(formula=["TRUE"],fill=fill("DDEBF7")))
for col in "FJ":
    rg=f"{col}{D0}:{col}{DL}"
    m.conditional_formatting.add(rg,FormulaRule(formula=[f'{col}{D0}="미달"'],fill=BAD,font=Font(name=FN,bold=True,color="9C0006")))
    m.conditional_formatting.add(rg,FormulaRule(formula=[f'{col}{D0}="주의"'],fill=WARN,font=Font(name=FN,bold=True,color="9C5700")))
    m.conditional_formatting.add(rg,FormulaRule(formula=[f'{col}{D0}="적정"'],fill=GOOD,font=Font(name=FN,bold=True,color="006100")))
# --- 4. 역산
Q0=DL+3
step(m,f"A{Q0}","4단계 · 거꾸로 계산하기 — 목표 확보율을 지키려면 재학생을 최대 몇 명까지 늘릴 수 있나?","K"+str(Q0))
note(m,f"A{Q0+1}","목표 확보율(노란 칸)은 바꿀 수 있습니다. 결과가 음수(−)이면 그 목표를 맞추려면 오히려 그만큼 재학생이 줄어야(또는 면적이 늘어야) 한다는 뜻입니다.","K"+str(Q0+1),size=9)
header(m,Q0+2,"ABCDEFGH",["목표 확보율\n(%)","허용 최대\n재학생 기준면적(㎡)","현재보다 더 늘릴 수\n있는 기준면적(㎡)","","인문사회만 늘릴 때\n최대 증가 인원","","자연·공학·예체능만\n늘릴 때 최대 증가 인원",""],height=36)
m.merge_cells(f"C{Q0+2}:D{Q0+2}"); m.merge_cells(f"E{Q0+2}:F{Q0+2}"); m.merge_cells(f"G{Q0+2}:H{Q0+2}")
targets=[130,125,120,115,110,105,100]
for i,t in enumerate(targets):
    r=Q0+3+i
    put(m,f"A{r}",t,F(True,BLUE,11),YEL,fmt='0"%"',al=C)
    put(m,f"B{r}",f"=$E$7/A{r}*100",fmt=AREA)
    put(m,f"C{r}",f"=B{r}-$E$6",fmt=AREA); m.merge_cells(f"C{r}:D{r}")
    put(m,f"E{r}",f"=IF(C{r}>=0,ROUNDDOWN(C{r}/12,0),-ROUNDUP(-C{r}/12,0))",F(True),fmt='+#,##0"명";-#,##0"명";0"명"',al=C); m.merge_cells(f"E{r}:F{r}")
    put(m,f"G{r}",f"=IF(C{r}>=0,ROUNDDOWN(C{r}/14,0),-ROUNDUP(-C{r}/14,0))",F(True),fmt='+#,##0"명";-#,##0"명";0"명"',al=C); m.merge_cells(f"G{r}:H{r}")
QL=Q0+3+len(targets)-1
for rg in (f"E{Q0+3}:H{QL}",):
    m.conditional_formatting.add(rg,CellIsRule(operator="lessThan",formula=["0"],font=Font(name=FN,bold=True,color="C00000")))
# --- 5. 계산기
K1=QL+3
step(m,f"A{K1}","5단계 · 내 시나리오 계산기 — 여러 계열을 섞어서 바꿔 보기","K"+str(K1))
note(m,f"A{K1+1}","예) 대학원 공학 +80명, 학부 인문사회 −150명, 기숙사 신축 +2,000㎡ 처럼 동시에 넣어 보세요. (학부·대학원 구분 없이 계열별로 합쳐서 입력)","K"+str(K1+1),size=9)
header(m,K1+2,"ABCD",["항목","증감 입력","1인당\n기준면적","기준면적 변화(㎡)"],height=32)
calc=[("인문사회 재학생(명)",-150,12),("자연과학 재학생(명)",0,14),("공학 재학생(명)",80,14),("예체능 재학생(명)",0,14)]
for i,(lab,v,u) in enumerate(calc):
    r=K1+3+i
    put(m,f"A{r}",lab,F(True),SUB,al=C); put(m,f"B{r}",v,F(True,BLUE,11),YEL,fmt='+#,##0;-#,##0;0',al=C)
    put(m,f"C{r}",u,fmt='0"㎡"',al=C); put(m,f"D{r}",f"=B{r}*C{r}",fmt='+#,##0;-#,##0;"-"')
ra=K1+7
put(m,f"A{ra}","교사 인정면적 증감(㎡)",F(True),SUB,al=C); put(m,f"B{ra}",2000,F(True,BLUE,11),YEL,fmt='+#,##0;-#,##0;0',al=C)
put(m,f"C{ra}","",); put(m,f"D{ra}","신축 +, 철거 −",F(size=9,color=GRAY),al=C)
# outputs at F..K
outs=[("재학생 증감 합계(명)",f"=SUM(B{K1+3}:B{K1+6})",'+#,##0;-#,##0;0'),
      ("새 재학생 수(명)",f"=E5+F{K1+3}",CNT),
      ("새 재학생 기준면적(㎡)",f"=E6+SUM(D{K1+3}:D{K1+6})",AREA),
      ("새 교사 인정면적(㎡)",f"=E7+B{ra}",AREA),
      ("새 재학생기준 확보율",f"=IF(J{K1+5}<=0,0,J{K1+6}/J{K1+5}*100)",PCT),
      ("현재 대비 변화",f"=J{K1+7}-J5",PP)]
put(m,f"F{K1+2}","계산 결과",F(True,"FFFFFF"),HDR,al=C); m.merge_cells(f"F{K1+2}:K{K1+2}")
for i,(lab,fm,fmt) in enumerate(outs):
    r=K1+3+i
    put(m,f"F{r}",lab,F(True),SUB,al=C); m.merge_cells(f"F{r}:I{r}")
    big=i in (4,5)
    put(m,f"J{r}",fm,F(True,NAVY,13 if big else 11),KPI,fmt=fmt,al=C); m.merge_cells(f"J{r}:K{r}")
for i in range(6): pass
# fix references: outputs rows K1+3..K1+8 at J; note J{K1+3}=sum, J{K1+5}=new B, J{K1+6}=new area
m[f"J{K1+4}"]=f"=E5+J{K1+3}"
put(m,f"A{K1+10}",f'="▶ 결과: 재학생 "&TEXT(J{K1+3},"+#,##0;-#,##0;0")&"명, 면적 "&TEXT(B{ra},"+#,##0;-#,##0;0")&"㎡ 변동 시 재학생기준 확보율은 "&TEXT(J5,"0.0")&"% → "&TEXT(J{K1+7},"0.0")&"% ("&TEXT(J{K1+8},"+0.0;-0.0;0.0")&"%p)로, 판정은 「"&IF(J{K1+7}<100,"미달",IF(J{K1+7}<{TH},"주의","적정"))&"」입니다."',
    F(True,"C00000",11),fill("FFF2F2"),al=L)
m.merge_cells(f"A{K1+10}:K{K1+10}"); m.row_dimensions[K1+10].height=28
note(m,f"A{K1+12}","※ 이 시트는 정보공시 단위(천안+서울 합산) 기준입니다. 입학정원기준 확보율은 재학생 수와 관계없이 입학정원·면적이 바뀔 때만 변합니다. 연도별 공식 예측은 입력 시트에 반영하세요.","K"+str(K1+12),size=9)
# chart
ch2=LineChart(); ch2.title="재학생 증감에 따른 재학생기준 확보율"; ch2.height=9; ch2.width=15
ch2.add_data(Reference(m,min_col=4,min_row=D0-1,max_row=DL),titles_from_data=True)
ch2.add_data(Reference(m,min_col=8,min_row=D0-1,max_row=DL),titles_from_data=True)
ch2.set_categories(Reference(m,min_col=1,min_row=D0,max_row=DL)); ch2.y_axis.title="확보율(%)"; ch2.x_axis.title="재학생 증감(명)"
ch2.legend.position="b"; ch2.y_axis.majorGridlines=None; ch2.x_axis.delete=False; ch2.y_axis.delete=False
for sr,nm,cl in zip(ch2.series,["인문사회가 바뀔 때","자연·공학·예체능이 바뀔 때"],["2F5597","548235"]):
    sr.tx=SeriesLabel(v=nm); sr.graphicalProperties.line.solidFill=cl; sr.graphicalProperties.line.width=25000
ch2.y_axis.scaling.min=80
m.add_chart(ch2,f"L{S0}")
for col,w in zip("ABCDEFGHIJK",[14,12,13,19,12,9,13,11,11,9,9]): m.column_dimensions[col].width=w
m.freeze_panes="A4"

# ============ 간편계산기 ============
q=quick; TH="사용안내!$C$27"
def judge(x): return f'=IF(NOT(ISNUMBER({x})),"",IF({x}<100,"미달",IF({x}<{TH},"주의","적정")))'
def judge_cf(rng,top):
    for t,fl_,fc in (("미달",BAD,"9C0006"),("주의",WARN,"9C5700"),("적정",GOOD,"006100")):
        q.conditional_formatting.add(rng,FormulaRule(formula=[f'{top}="{t}"'],fill=fl_,font=Font(name=FN,bold=True,color=fc)))
# 숨김 기준값 (J열 값, K열 설명) — 모두 2025 실적에서 자동으로 가져옴
hv=[("J2","=기초자료_2025!$E$16","교사 인정면적 2025(㎡)",AREA2),
    ("J3","=기초자료_2025!$C$39","재학생 기준면적 2025(㎡)",AREA),
    ("J4","=SUM(기초자료_2025!$G$25:$G$27)","재학생 수 2025(명)",CNT),
    ("J5","=J3/J4","1인당 평균 기준면적(㎡) = J3/J4","0.0000"),
    ("J6","=J2/J3*100","재학생기준 확보율 2025(%)","0.000")]
for ref,fm,lab,fmt in hv:
    put(q,ref,fm,fmt=fmt,border=False); put(q,"K"+ref[1:],lab,F(color=GRAY),border=False)
banner(q,"A1","간편 계산기 · 재학생 수만 넣으면 예상 교사확보율이 나옵니다","H1")
note(q,"A2","노란 칸에 예상 재학생 수만 넣으세요. 나머지는 자동 계산됩니다. 처음에는 2025년 실적이 들어 있어 2025 공시값과 같은 결과가 나옵니다.","H2",color="C00000")
put(q,"A3",'="※ 건물(교사) 인정면적은 2025년 실적 "&TEXT($J$2,"#,##0")&"㎡ 그대로라고 가정합니다. 신축·철거를 반영하거나 연도별로 예측하려면 \'입력\' 시트를 쓰세요."',F(color=GRAY,size=9),al=L,border=False); q.merge_cells("A3:H3")
# --- 방법 ①
step(q,"A5","방법 ①  전체 재학생 수만 넣기 — 가장 간단","H5")
for rng,lab in (("A6:C6","예상 재학생 수 (명)  ✎ 입력"),("E6:F6","예상 교사확보율"),("G6:H6","판정")):
    put(q,rng.split(":")[0],lab,F(True,"FFFFFF",11),fill("2F5597"),al=C); q.merge_cells(rng)
put(q,"A7","=$J$4",F(True,BLUE,22),YEL,fmt=CNT,al=C); q.merge_cells("A7:C8")
put(q,"D7","▶",F(True,NAVY,22),al=C,border=False); q.merge_cells("D7:D8")
put(q,"E7",'=IF(AND(ISNUMBER(A7),A7>0),$J$2/(A7*$J$5)*100,"")',F(True,NAVY,22),KPI,fmt=PCT,al=C); q.merge_cells("E7:F8")
put(q,"G7",judge("E7"),F(True,NAVY,22),KPI,al=C); q.merge_cells("G7:H8")
judge_cf("G7","$G$7")
q.row_dimensions[7].height=26; q.row_dimensions[8].height=20
put(q,"A9",'=IF(E7="","재학생 수를 0보다 큰 숫자로 넣어 주세요.","재학생이 "&TEXT(A7,"#,##0")&"명이면 예상 교사확보율은 "&TEXT(E7,"0.0")&"%입니다. (2025년 실적 "&TEXT($J$6,"0.0")&"%보다 "&TEXT(ROUND(E7-$J$6,1),"+0.0;-0.0;0.0")&"%p, "&IF(E7>=100,"법정 기준 100% 충족","법정 기준 100% 미달")&")")',
    F(True,NAVY,11),al=L,border=False); q.merge_cells("A9:H9"); q.row_dimensions[9].height=22
put(q,"A10",'=IF(E7="","","계산: 교사 인정면적 "&TEXT($J$2,"#,##0")&"㎡ ÷ (재학생 "&TEXT(A7,"#,##0")&"명 × 1인당 평균 "&TEXT($J$5,"0.00")&"㎡) × 100 = "&TEXT(E7,"0.0")&"%")',
    F(color="404040"),al=L,border=False); q.merge_cells("A10:H10")
put(q,"A11",'="※ 1인당 평균 "&TEXT($J$5,"0.00")&"㎡는 2025년 계열 구성(인문사회 1인당 12㎡, 자연과학·공학·예체능 14㎡)이 그대로 유지된다고 보고 계산한 값입니다. 계열 구성이 크게 바뀔 것 같으면 아래 방법 ②를 쓰세요."',
    F(color=GRAY,size=9),al=L,border=False); q.merge_cells("A11:H11"); q.row_dimensions[11].height=28
# --- 방법 ②
step(q,"A13","방법 ②  계열별 재학생 수 넣기 — 더 정확 (천안·서울, 학부·대학원 모두 합친 인원)","H13")
header(q,14,"ABCDE",["계열","1인당\n기준면적","2025 실적\n(참고)","예상 재학생 수\n✎ 입력","필요 기준면적(㎡)\n= 인원 × 1인당"],height=36)
for i,(nm,bcol,u) in enumerate(flds):
    r=15+i
    put(q,f"A{r}",nm,F(True),al=C); put(q,f"B{r}",u,fmt='0"㎡"',al=C)
    put(q,f"C{r}",f"=SUM(기초자료_2025!{bcol}25:{bcol}27)",F(color=GREEN),fill("F2F2F2"),fmt=CNT,al=C)
    put(q,f"D{r}",f"=C{r}",F(True,BLUE,12),YEL,fmt=CNT,al=C)
    put(q,f"E{r}",f"=B{r}*D{r}",fmt=AREA)
put(q,"A19","합계",F(True),TOT,al=C); put(q,"B19","",fl=TOT)
for col in "CDE": put(q,f"{col}19",f"=SUM({col}15:{col}18)",F(True),TOT,fmt=AREA if col=="E" else CNT,al=None if col=="E" else C)
put(q,"G14","예상 교사확보율",F(True,"FFFFFF",11),fill("2F5597"),al=C); q.merge_cells("G14:H14")
put(q,"G15",'=IF(E19>0,$J$2/E19*100,"")',F(True,NAVY,22),KPI,fmt=PCT,al=C); q.merge_cells("G15:H16")
put(q,"G17","판정",F(True,"FFFFFF",11),fill("2F5597"),al=C); q.merge_cells("G17:H17")
put(q,"G18",judge("G15"),F(True,NAVY,18),KPI,al=C); q.merge_cells("G18:H19")
judge_cf("G18","$G$18")
put(q,"A20",'=IF(G15="","재학생 수를 입력해 주세요.","재학생 "&TEXT(D19,"#,##0")&"명(인문사회 "&TEXT(D15,"#,##0")&" · 자연과학 "&TEXT(D16,"#,##0")&" · 공학 "&TEXT(D17,"#,##0")&" · 예체능 "&TEXT(D18,"#,##0")&")이면 예상 교사확보율은 "&TEXT(G15,"0.0")&"%입니다. (2025년 실적 대비 "&TEXT(ROUND(G15-$J$6,1),"+0.0;-0.0;0.0")&"%p)")',
    F(True,NAVY,11),al=L,border=False); q.merge_cells("A20:H20"); q.row_dimensions[20].height=30
put(q,"A21",'=IF(G15="","","계산: 교사 인정면적 "&TEXT($J$2,"#,##0")&"㎡ ÷ 필요 기준면적 합계 "&TEXT(E19,"#,##0")&"㎡ × 100 = "&TEXT(G15,"0.0")&"%")',
    F(color="404040"),al=L,border=False); q.merge_cells("A21:H21")
# --- 한눈에 보는 표
step(q,"A23","한눈에 보는 표 — 전체 재학생 수별 예상 교사확보율 (방법 ① 기준)","H23")
put(q,"A24",f'="▶ 재학생이 "&TEXT(ROUNDDOWN($J$2/$J$5,0),"#,##0")&"명 이하면 법정 기준(100%) 충족, "&TEXT(ROUNDDOWN($J$2/($J$5*{TH}/100),0),"#,##0")&"명 이하면 주의 기준("&{TH}&"%) 이상입니다."',
    F(True,"C00000",11),fill("FFF2F2"),al=L); q.merge_cells("A24:H24"); q.row_dimensions[24].height=22
header(q,25,"ABCD",["재학생 수(명)","예상 확보율","2025 실적 대비","판정"],height=26)
QT0=26; qn=list(range(11000,18001,500))
for i,n in enumerate(qn):
    r=QT0+i
    put(q,f"A{r}",n,F(True),fmt=CNT,al=C)
    put(q,f"B{r}",f"=$J$2/(A{r}*$J$5)*100",F(True),fmt=PCT,al=C)
    put(q,f"C{r}",f"=ROUND(B{r}-$J$6,1)",fmt=PP,al=C)
    put(q,f"D{r}",judge(f"B{r}"),F(True),al=C)
QTL=QT0+len(qn)-1
judge_cf(f"D{QT0}:D{QTL}",f"D{QT0}")
note(q,f"A{QTL+2}","※ 판정: 100% 미만 = 미달, 100% 이상 ~ 주의 기준 미만 = 주의, 그 이상 = 적정 (주의 기준은 사용안내 시트에서 변경). 이 계산기는 재학생기준 확보율만 계산합니다. 입학정원기준 확보율은 재학생 수와 관계없이 입학정원·면적이 바뀔 때만 변합니다.","H"+str(QTL+2),size=9)
q.row_dimensions[QTL+2].height=30
for col,w in zip("ABCDEFGH",[16,12,14,16,16,4,14,14]): q.column_dimensions[col].width=w
q.column_dimensions["J"].hidden=True; q.column_dimensions["K"].hidden=True

# ============ 사용안내 ============
g=guide
banner(g,"A1","교사(校舍)시설 확보율 예측 시뮬레이션 — 사용안내","H1")
note(g,"A2","2025 시설현황 집계표와 2025 정보공시(14-라)를 바탕으로, 재학생 수가 바뀌면 교사확보율이 어떻게 변하는지 미리 계산해 보는 파일입니다.","H2")
step(g,"A4","이렇게 쓰세요 (3단계)  ※ 재학생 수만으로 빠르게 보려면 첫 번째 '간편계산기' 시트만 쓰면 됩니다","H4")
steps=[("1","입력 시트","노란 칸에 연도별 예상 재학생 수를 넣습니다. 정원 조정이나 건물 신축·철거가 있으면 STEP 2·3도 고칩니다."),
       ("2","결과요약 시트","맨 위 '조회 연도'를 고르면 그 해 확보율·여유면적·판정과 쉬운 설명 문장이 나옵니다. 아래 표와 그래프로 연도별 흐름을 봅니다."),
       ("3","민감도 시트","재학생 증감별 확보율 변화, 목표 확보율을 지키는 최대 증가 인원, 계열을 섞은 시나리오 계산을 단계별로 확인합니다.")]
for i,(n,sh,t) in enumerate(steps):
    row=5+i
    put(g,f"A{row}",n,F(True,"FFFFFF",14),fill("2F5597"),al=C)
    put(g,f"B{row}",sh,F(True,NAVY),SUB,al=C); g.merge_cells(f"B{row}:C{row}")
    put(g,f"D{row}",t,al=L); g.merge_cells(f"D{row}:H{row}"); g.row_dimensions[row].height=34
step(g,"A9","칸 색깔의 의미","H9")
legend=[(YEL,"노란 칸","직접 고치는 입력 칸"),(fill("F2F2F2"),"회색 칸","2025 실적(기초자료와 연결, 수정 불필요)"),(None,"흰 칸","자동 계산 칸 (수정 금지)")]
for i,(fl,a,t) in enumerate(legend):
    row=10+i; put(g,f"A{row}","",fl=fl); put(g,f"B{row}",a,F(True),al=C); g.merge_cells(f"B{row}:C{row}")
    put(g,f"D{row}",t,al=L); g.merge_cells(f"D{row}:H{row}")
step(g,"A14","용어 쉽게 풀기","H14")
terms=[("교사확보율","확보율 계산에 인정되는 건물 면적 ÷ 법정 기준면적 × 100. 100% 이상이면 기준 충족"),
       ("재학생 기준면적(B)","계열별 재학생 수 × 1인당 기준면적(인문사회 12㎡, 자연과학·공학·예체능 14㎡)"),
       ("입학정원 기준면적(A)","계열별 입학정원 × 수업연한 × 1인당 기준면적. 재학생 수가 바뀌어도 변하지 않음"),
       ("교사 인정면적(C+D+E)","기본시설(강의실·실험실·도서관 등) + 지원시설(기숙사·강당 등) + 연구시설. 부속·기타시설은 제외"),
       ("여유 면적","교사 인정면적 − 재학생 기준면적. 이 면적을 1인당 기준면적으로 나누면 더 받을 수 있는 재학생 수"),
       ("법정 의무 기준","재학생이 입학정원보다 적으면 재학생 기준면적 이상, 그렇지 않으면 입학정원 기준면적 이상 보유해야 함")]
for i,(a,t) in enumerate(terms):
    row=15+i; put(g,f"A{row}",a,F(True),SUB,al=C); g.merge_cells(f"A{row}:C{row}")
    put(g,f"D{row}",t,al=L); g.merge_cells(f"D{row}:H{row}"); g.row_dimensions[row].height=30
step(g,"A22","실적 현황","C22")
for col,v in zip("DEF",["2024 실적","2025 실적","변화"]): put(g,f"{col}22",v,F(True,"FFFFFF"),HDR,al=C)
for i,(lab,cl,fmt,dfmt) in enumerate([("재학생기준 확보율","F",PCT,PP),("입학정원기준 확보율","G",PCT,PP),("여유 면적(㎡)","J",AREA,'+#,##0;-#,##0;"-"')]):
    r=23+i; put(g,f"A{r}",lab,F(True),SUB,al=C); g.merge_cells(f"A{r}:C{r}")
    put(g,f"D{r}",f"=결과요약!{cl}{T0}",F(True,NAVY,12),fmt=fmt,al=C)
    put(g,f"E{r}",f"=결과요약!{cl}{T0+1}",F(True,NAVY,12),fmt=fmt,al=C)
    put(g,f"F{r}",f"=E{r}-D{r}",F(True,"C00000",11),fmt=dfmt,al=C)
step(g,"A26","설정","H26")
put(g,"A27","'주의' 판정 기준(%)",F(True),SUB,al=C); g.merge_cells("A27:B27")
put(g,"C27",110,F(True,BLUE,12),YEL,fmt="0",al=C)
put(g,"D27","확보율이 100% 이상이지만 이 값보다 낮으면 '주의'로 표시합니다. 내부 관리용 임의 기준이므로 필요에 맞게 바꾸세요.",F(size=9,color=GRAY),al=L); g.merge_cells("D27:H27")
g.row_dimensions[27].height=30
for col,w in zip("ABCDEFGH",[8,12,12,16,16,16,16,16]): g.column_dimensions[col].width=w
g.column_dimensions["A"].width=10
for ws in wb:
    ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=9
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
for ws,r in ((quick,2),(inp,2),(guide,2),(base25,2),(base24,2),(sen,2)): ws.row_dimensions[r].height=30
inp.row_dimensions[7].height=28; inp.row_dimensions[26].height=28
wb.active=0
wb.save("교사시설확보율_예측시뮬레이션.xlsx")
